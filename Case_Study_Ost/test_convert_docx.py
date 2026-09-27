import zipfile
import re

z = zipfile.ZipFile('CSR.docx', 'r')
with zipfile.ZipFile('CSR_standard.docx', 'w', zipfile.ZIP_DEFLATED) as out_z:
    for item in z.infolist():
        data = z.read(item.filename)
        if item.filename.endswith(('.xml', '.rels')):
            data = data.replace(b'http://purl.oclc.org/ooxml/officeDocument/relationships', b'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
            data = data.replace(b'http://purl.oclc.org/ooxml/package/relationships', b'http://schemas.openxmlformats.org/package/2006/relationships')
            data = data.replace(b'http://purl.oclc.org/ooxml/wordprocessingml/main', b'http://schemas.openxmlformats.org/wordprocessingml/2006/main')
            data = data.replace(b'http://purl.oclc.org/ooxml/drawingml/main', b'http://schemas.openxmlformats.org/drawingml/2006/main')
            data = data.replace(b'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing', b'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing')
            data = data.replace(b'http://purl.oclc.org/ooxml/officeDocument/math', b'http://schemas.openxmlformats.org/officeDocument/2006/math')
        out_z.writestr(item, data)

print("Converted to standard namespaces!")

import docx
try:
    doc = docx.Document('CSR_standard.docx')
    print("python-docx successfully opened CSR_standard.docx!")
    print("Paragraphs:", len(doc.paragraphs))
    print("Tables:", len(doc.tables))
    print("Sections:", len(doc.sections))
except Exception as e:
    print("Failed:", type(e), e)
