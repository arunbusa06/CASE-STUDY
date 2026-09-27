import docx

doc = docx.Document('CSR_standard.docx')
for i in range(75, 175):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        if p.text.strip():
            print(f"P{i} ({p.style.name}): {p.text.strip()[:90]}")
