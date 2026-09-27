import os
import docx

from frontmatter_updater import update_complete_frontmatter
from build_section_intro import generate_section_intro
from build_section_analysis import generate_section_analysis
from build_section_arch_api import generate_section_arch_api
from build_section_testing import generate_section_testing
from build_section_security_tech import generate_section_security_tech
from build_section_sdg_challenges import generate_section_sdg_challenges

def main():
    print("=== STARTING CASE STUDY GENERATION ===")
    
    # 1. Load clean template
    doc = docx.Document('CSR_clean.docx')
    print(f"Loaded CSR_clean.docx with {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} sections.")
    
    # 2. Update Frontmatter (Section 0 and Section 1)
    update_complete_frontmatter(doc)
    
    # 3. Remove old body elements from 176 to len(body)-1
    body = doc._element.body
    print(f"Body length before removal: {len(body)}")
    old_elements = list(body[176:-1])
    for el in old_elements:
        body.remove(el)
    print(f"Body length after removal: {len(body)}")
    
    # 4. Generate all Sections into Section 4
    generate_section_intro(doc)
    generate_section_analysis(doc)
    generate_section_arch_api(doc)
    generate_section_testing(doc)
    generate_section_security_tech(doc)
    generate_section_sdg_challenges(doc)
    
    # 5. Save the generated document
    output_filename = "Food_Delivery_API_Testing_Case_Study.docx"
    doc.save(output_filename)
    print(f"Saved generated document to {output_filename} successfully!")
    
    # Overwrite CSR.docx as well so the user's primary document is updated
    doc.save("CSR.docx")
    print("Updated CSR.docx successfully!")
    
    # 6. Verification
    v_doc = docx.Document(output_filename)
    print("=== VERIFICATION SUMMARY ===")
    print(f"Total Paragraphs: {len(v_doc.paragraphs)}")
    print(f"Total Tables: {len(v_doc.tables)}")
    print(f"Total Sections: {len(v_doc.sections)}")
    print(f"Section 4 Column Count: {v_doc.sections[-1]._sectPr.xpath('./w:cols/@w:num')}")
    print("Verification complete!")

if __name__ == '__main__':
    main()
