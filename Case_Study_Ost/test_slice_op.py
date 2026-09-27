import docx

doc = docx.Document('CSR_clean.docx')
body = doc._element.body

# Delete elements from 176 to len(body)-1
for el in list(body[176:-1]):
    body.remove(el)

print(f"Body length after removal: {len(body)}")
# Add a test paragraph
p = doc.add_paragraph("This is a test paragraph in Section 4.")
# Save
doc.save('test_sliced.docx')

# Reopen
doc2 = docx.Document('test_sliced.docx')
print(f"Reopened successfully! Paragraphs: {len(doc2.paragraphs)}, Sections: {len(doc2.sections)}")
print(f"Last section columns: {doc2.sections[-1]._sectPr.xpath('./w:cols/@w:num')}")
