import docx

doc = docx.Document('Food_Delivery_API_Testing_Case_Study.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")

word_count = sum(len(p.text.split()) for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            word_count += len(cell.text.split())
print(f"Total Document Word Count: {word_count} words")

print("\n--- Document Outline / Major Headings Found ---")
for p in doc.paragraphs:
    txt = p.text.strip()
    if any(txt.startswith(x) for x in ['I. ', 'II. ', 'III. ', 'IV. ', 'V. ', 'VI. ', 'VII. ', 'VIII. ', 'IX. ', 'X. ', 'XI. ', 'XII. ', 'REFERENCES', 'APPENDIX']):
        print(f"  {txt}")

print("\n--- Tables Summary ---")
for i, t in enumerate(doc.tables):
    first_cell = t.cell(0, 0).text.strip().replace('\n', ' ')[:35]
    print(f"  Table {i}: {len(t.rows)} rows x {len(t.columns)} cols | First Cell: '{first_cell}'")

print("\n--- Checking Embedded Figures ---")
body_xml = doc._element.body.xml
blip_count = body_xml.count('embed=')
print(f"Embedded blip references in body XML: {blip_count}")
