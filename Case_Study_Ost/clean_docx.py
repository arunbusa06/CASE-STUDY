import zipfile
import re

def pt_to_twips(match):
    prefix = match.group(1)
    val = float(match.group(2))
    twips = int(round(val * 20))
    return f'{prefix}="{twips}"'

z = zipfile.ZipFile('CSR.docx', 'r')
with zipfile.ZipFile('CSR_clean.docx', 'w', zipfile.ZIP_DEFLATED) as out_z:
    for item in z.infolist():
        data = z.read(item.filename)
        if item.filename.endswith(('.xml', '.rels')):
            # Fix namespaces
            data = data.replace(b'http://purl.oclc.org/ooxml/officeDocument/relationships', b'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
            data = data.replace(b'http://purl.oclc.org/ooxml/package/relationships', b'http://schemas.openxmlformats.org/package/2006/relationships')
            data = data.replace(b'http://purl.oclc.org/ooxml/wordprocessingml/main', b'http://schemas.openxmlformats.org/wordprocessingml/2006/main')
            data = data.replace(b'http://purl.oclc.org/ooxml/drawingml/main', b'http://schemas.openxmlformats.org/drawingml/2006/main')
            data = data.replace(b'http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing', b'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing')
            data = data.replace(b'http://purl.oclc.org/ooxml/officeDocument/math', b'http://schemas.openxmlformats.org/officeDocument/2006/math')
            
            # Convert pt to twips in attributes
            text = data.decode('utf-8', errors='ignore')
            text = re.sub(r'(w:\w+)=[\x22\x27]([\d\.]+)pt[\x22\x27]', pt_to_twips, text)
            data = text.encode('utf-8')
        out_z.writestr(item, data)

print("Converted CSR.docx to standard CSR_clean.docx!")

import docx
doc = docx.Document('CSR_clean.docx')
print("Successfully opened CSR_clean.docx in python-docx!")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Sections: {len(doc.sections)}")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: page_width={s.page_width.pt}pt, page_height={s.page_height.pt}pt, top={s.top_margin.pt}pt")
for i, t in enumerate(doc.tables):
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols, width={t.cell(0, 0).width.pt}pt")
