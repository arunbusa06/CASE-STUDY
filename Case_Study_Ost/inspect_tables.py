import docx

doc = docx.Document('CSR_standard.docx')
for i in range(4, len(doc.tables)):
    tbl = doc.tables[i]
    print(f"\n--- Table {i} ---")
    print(f"Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}, Style: {tbl.style.name if tbl.style else 'None'}")
    first_row = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    print(f"Header: {first_row}")
    # inspect widths
    col_widths = [tbl.cell(0, c).width.pt if tbl.cell(0, c).width else 'None' for c in range(len(tbl.columns))]
    print(f"Widths: {col_widths}")
