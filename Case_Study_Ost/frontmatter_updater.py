import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_complete_frontmatter(doc):
    print("Updating frontmatter and outline cleanly...")
    
    # Title Page
    doc.paragraphs[1].text = "Food Delivery Application: API Testing and Error Handling for Web Services"
    doc.paragraphs[3].text = "[92400120478] Rajvirsinh Dodiya"
    doc.paragraphs[4].text = "[92400120494] Harsh Kanjariya"
    doc.paragraphs[5].text = "[92400120614] Arun Busa"
    
    doc.paragraphs[8].text = "COMPUTER ENGINEERING"
    doc.paragraphs[11].text = "2026–2027"
    doc.paragraphs[14].text = "5TH"
    doc.paragraphs[17].text = "OPEN SOURCE TECHNOLOGIES (01CE0526)"
    doc.paragraphs[20].text = "5EV4"
    doc.paragraphs[23].text = "A"
    
    # Table 1: Evaluation Table
    doc.tables[1].cell(1, 1).text = "Food Delivery Application: API Testing and Error Handling for Web Services"
    doc.tables[1].cell(1, 2).text = "CO5: API Testing & Software Reliability"
    
    # Course / CO / SDG block
    doc.paragraphs[42].text = "Food Delivery Application: API Testing and Error Handling for Web Services"
    doc.paragraphs[44].text = "Course Name and Course Code: OPEN SOURCE TECHNOLOGIES (01CE0526)"
    doc.paragraphs[46].text = "COs Mapped – CO5: API Testing & Software Reliability"
    doc.paragraphs[48].text = "SDGs Mapped - SDG 9 – Industry, Innovation and Infrastructure"
    doc.paragraphs[49].text = ""
    
    # Clear ALL paragraphs from 54 to 166
    for p_idx in range(54, 167):
        doc.paragraphs[p_idx].text = ""

    # Outline updating (P54 onwards)
    outline_items = [
        "I. Introduction",
        "II. Application / System Background",
        "III. Problem Analysis",
        "IV. Proposed API Testing and Error-Handling Solution",
        "V. System Architecture and Technology Stack",
        "VI. API and Database Analysis",
        "VII. Software Quality and API Testing",
        "VIII. API Security and Reliability Analysis",
        "IX. Technical Analysis",
        "X. SDG Mapping and Justification",
        "XI. Challenges and Recommendations",
        "XII. Conclusion",
        "References",
        "Appendix"
    ]
    for idx, item in enumerate(outline_items):
        p_target = 54 + idx
        doc.paragraphs[p_target].text = item
        doc.paragraphs[p_target].paragraph_format.space_before = Pt(1)
        doc.paragraphs[p_target].paragraph_format.space_after = Pt(2)

    # Paper Title (P167)
    doc.paragraphs[167].text = "Food Delivery Application: API Testing and Error Handling for Web Services"
    
    # Author Table (Table 3)
    auth_tbl = doc.tables[3]
    auth_tbl.cell(0, 0).text = "Rajvirsinh Dodiya [92400120478]\nDepartment of Computer Engineering\nMarwadi University\nRajkot, Gujarat, India\nrajvirsinh.dodiya120478@marwadiuniversity.ac.in"
    auth_tbl.cell(0, 1).text = "Harsh Kanjariya [92400120494]\nDepartment of Computer Engineering\nMarwadi University\nRajkot, Gujarat, India\nharsh.kanjariya120494@marwadiuniversity.ac.in"
    auth_tbl.cell(0, 2).text = "Arun Busa [92400120614]\nDepartment of Computer Engineering\nMarwadi University\nRajkot, Gujarat, India\narun.busa120614@marwadiuniversity.ac.in"
    
    for c_idx in range(3):
        cell = auth_tbl.cell(0, c_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)

    print("Frontmatter and outline cleanly updated!")
